"""
Script de diagnostic : pourquoi un .docx extrait si peu de texte ?
--------------------------------------------------------------------
Analyse la structure interne d'un fichier Word pour détecter :
- le nombre de paragraphes avec du texte réel
- le nombre d'images intégrées
- le nombre de zones de texte / formes (souvent ignorées par python-docx)
- le nombre de tableaux

Usage :
    python diagnostic_docx.py "data/docx/QES24019 Consignes d-evacuation AVL Maroc.docx"
"""

import sys
import zipfile
from pathlib import Path

from docx import Document as DocxDocument


def diagnostiquer(chemin_fichier: str):
    chemin = Path(chemin_fichier)
    print(f"📄 Analyse de : {chemin.name}\n")

    # 1. Ce que python-docx voit (paragraphes + tableaux)
    doc = DocxDocument(chemin)
    paragraphes_non_vides = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"Paragraphes avec texte : {len(paragraphes_non_vides)}")
    for p in paragraphes_non_vides:
        print(f"   - {p[:80]}")

    print(f"\nTableaux détectés : {len(doc.tables)}")
    for i, table in enumerate(doc.tables):
        print(f"   Tableau {i+1} : {len(table.rows)} lignes x {len(table.columns)} colonnes")

    # 2. Ce qu'il y a dans le fichier .docx (qui est en réalité un zip)
    print("\n--- Contenu brut du fichier (structure interne) ---")
    with zipfile.ZipFile(chemin) as z:
        noms = z.namelist()

        images = [n for n in noms if n.startswith("word/media/")]
        print(f"Images intégrées : {len(images)}")
        for img in images:
            print(f"   - {img}")

        # Recherche de zones de texte / formes dans le XML brut
        with z.open("word/document.xml") as f:
            xml_content = f.read().decode("utf-8", errors="ignore")

        nb_textbox = xml_content.count("<w:txbxContent")
        nb_drawing = xml_content.count("<w:drawing")
        nb_pict = xml_content.count("<w:pict")

        print(f"\nZones de texte (txbxContent) : {nb_textbox}")
        print(f"Éléments 'drawing' (images/formes/SmartArt) : {nb_drawing}")
        print(f"Éléments 'pict' (anciennes images) : {nb_pict}")

    print("\n--- Conclusion ---")
    if images or nb_drawing > 0:
        print("➡️  Ce document contient des IMAGES. Si l'info de sécurité est dedans")
        print("    (plan d'évacuation, pictogrammes avec texte), il faudra de l'OCR")
        print("    sur ces images pour les récupérer.")
    if nb_textbox > 0:
        print("➡️  Ce document contient des ZONES DE TEXTE, que python-docx ignore")
        print("    par défaut. On peut les extraire avec une méthode dédiée.")
    if not images and nb_drawing == 0 and nb_textbox == 0:
        print("➡️  Pas d'image ni de zone de texte détectée. Le document est")
        print("    peut-être simplement très court, ou le texte est dans un format")
        print("    inhabituel (SmartArt complexe, en-tête/pied de page...).")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage : python diagnostic_docx.py <chemin_vers_fichier.docx>")
        sys.exit(1)
    diagnostiquer(sys.argv[1])