"""
Étape 3 : Extraction du texte
------------------------------
Ce script lit tous les documents présents dans le dossier `data/`
(PDF, Word, TXT) et extrait leur contenu en texte brut.

Chaque document devient un dictionnaire :
    {
        "source": "nom_du_fichier.docx",
        "text": "contenu extrait ici...",
        "type": "document" ou "formulaire",
        "chemin_complet": "/chemin/vers/data/formulaires/nom_du_fichier.docx"
    }

Le champ "type" permet de distinguer :
  - "document"   : contenu informatif (politiques, règlement...)
  - "formulaire" : fichier à remplir/télécharger (les fichiers placés
                    dans data/formulaires/ sont automatiquement tagués ainsi)

Le champ "chemin_complet" permet, plus tard, de proposer le téléchargement
du fichier original directement dans l'interface Streamlit quand la réponse
s'appuie sur un formulaire.

Ces résultats seront utilisés à l'étape suivante (nettoyage + chunking).
"""

import os
import zipfile
from io import BytesIO
from pathlib import Path

import pdfplumber
import openpyxl
import pytesseract
from PIL import Image
from lxml import etree
from pdf2image import convert_from_path
from docx import Document as DocxDocument


# Namespace XML utilisé par Word pour le texte (nécessaire pour lire
# le contenu des zones de texte, que python-docx ignore par défaut)
NS_WORD = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


# Dossier racine des données (à adapter si besoin)
DATA_DIR = Path(__file__).resolve().parents[2] / "data"

# Sous-dossiers contenant du contenu informatif (politiques, règlement, FAQ...)
DOSSIERS_DOCUMENTS = ["pdf", "docx", "txt", "faq"]

# Sous-dossier contenant les formulaires à remplir/télécharger
DOSSIERS_FORMULAIRES = ["formulaires"]

# Jeux de langues à essayer pour l'OCR, du plus probable au moins probable.
# On évite de mélanger fra+ara par défaut : ça dégrade la reconnaissance
# quand un seul des deux est réellement présent sur la page.
CANDIDATS_LANGUES_OCR = ["fra+eng", "ara", "fra+ara+eng"]

# En dessous de ce nombre de caractères, on considère qu'un PDF n'a pas
# de texte exploitable et on bascule automatiquement sur l'OCR.
SEUIL_TEXTE_MINIMAL = 20

# Même logique pour les .docx : en dessous de ce seuil, on considère que
# le contenu utile est probablement dans des zones de texte ou des images
# (plans, infographies, pictogrammes) plutôt que dans des paragraphes normaux.
SEUIL_TEXTE_MINIMAL_DOCX = 100


def _score_qualite_ocr(texte: str) -> float:
    """
    Estime grossièrement la qualité d'un résultat OCR : proportion de
    caractères 'lettres' (plutôt que symboles/ponctuation parasite).
    Plus le score est proche de 1, plus le texte semble propre.
    """
    if not texte.strip():
        return 0.0
    nb_lettres = sum(1 for c in texte if c.isalpha())
    return nb_lettres / len(texte)


def extraire_texte_pdf_ocr(chemin_fichier: Path) -> str:
    """
    Extrait le texte d'un PDF scanné (image) via OCR (Tesseract).
    Convertit chaque page en image, puis teste plusieurs jeux de langues
    (français, arabe, mixte) et garde le résultat qui semble le plus propre,
    plutôt que de tout mélanger dans une seule passe (ce qui dégrade la
    qualité quand une seule langue est réellement présente).
    """
    pages_images = convert_from_path(chemin_fichier, dpi=300)

    meilleur_texte = ""
    meilleur_score = -1.0

    for langues in CANDIDATS_LANGUES_OCR:
        texte_pages = []
        for image in pages_images:
            texte_page = pytesseract.image_to_string(image, lang=langues)
            if texte_page.strip():
                texte_pages.append(texte_page)
        texte_candidat = "\n".join(texte_pages)

        score = _score_qualite_ocr(texte_candidat)
        if score > meilleur_score:
            meilleur_score = score
            meilleur_texte = texte_candidat

    return meilleur_texte


def extraire_texte_pdf(chemin_fichier: Path) -> str:
    """
    Extrait tout le texte d'un fichier PDF, page par page.

    Essaie d'abord l'extraction "classique" (texte natif du PDF, rapide).
    Si le résultat est vide ou trop court (PDF scanné = image), bascule
    automatiquement sur l'OCR (plus lent mais fonctionne sur les scans).
    """
    texte_complet = []
    with pdfplumber.open(chemin_fichier) as pdf:
        for page in pdf.pages:
            texte_page = page.extract_text()
            if texte_page:
                texte_complet.append(texte_page)

    texte = "\n".join(texte_complet)

    if len(texte.strip()) < SEUIL_TEXTE_MINIMAL:
        print(f"   ↳ Texte natif insuffisant, passage en OCR pour {chemin_fichier.name}...")
        texte = extraire_texte_pdf_ocr(chemin_fichier)

    return texte


def extraire_zones_de_texte_docx(chemin_fichier: Path) -> str:
    """
    Extrait le texte des zones de texte (text boxes) d'un .docx.

    python-docx ne lit que les paragraphes du corps du document : le texte
    placé dans des zones de texte, formes ou infographies (souvent utilisé
    pour les plans d'évacuation, schémas, etc.) est invisible pour lui.
    On va donc lire directement le XML interne du fichier Word pour
    retrouver tout le texte contenu dans les balises <w:txbxContent>.
    """
    morceaux = []
    with zipfile.ZipFile(chemin_fichier) as z:
        if "word/document.xml" not in z.namelist():
            return ""
        xml_bytes = z.read("word/document.xml")

    arbre = etree.fromstring(xml_bytes)

    # On cherche toutes les zones de texte, puis tout le texte à l'intérieur
    for zone_texte in arbre.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent"):
        textes_zone = zone_texte.findall(".//w:t", NS_WORD)
        contenu = " ".join(t.text for t in textes_zone if t.text)
        if contenu.strip():
            morceaux.append(contenu.strip())

    return "\n".join(morceaux)


def extraire_images_docx_ocr(chemin_fichier: Path) -> str:
    """
    Extrait le texte visible dans les images intégrées à un .docx via OCR.
    Utile pour les plans, pictogrammes ou schémas contenant du texte
    (ex: plan d'évacuation, consignes de sécurité illustrées).
    """
    morceaux = []
    with zipfile.ZipFile(chemin_fichier) as z:
        images = [n for n in z.namelist() if n.startswith("word/media/")]

        for nom_image in images:
            try:
                donnees_image = z.read(nom_image)
                image = Image.open(BytesIO(donnees_image)).convert("RGB")

                # On réutilise la même logique multi-langues que pour les PDF scannés
                meilleur_texte = ""
                meilleur_score = -1.0
                for langues in CANDIDATS_LANGUES_OCR:
                    texte_candidat = pytesseract.image_to_string(image, lang=langues)
                    score = _score_qualite_ocr(texte_candidat)
                    if score > meilleur_score:
                        meilleur_score = score
                        meilleur_texte = texte_candidat

                if meilleur_texte.strip():
                    morceaux.append(meilleur_texte.strip())
            except Exception:
                # Certaines images (logos, icônes vectorielles) ne sont pas
                # exploitables par OCR : on les ignore simplement.
                continue

    return "\n".join(morceaux)


def extraire_texte_docx(chemin_fichier: Path) -> str:
    """
    Extrait tout le texte d'un fichier Word (.docx), en combinant plusieurs
    sources de contenu :
      1. Paragraphes classiques du corps du document
      2. Tableaux
      3. Zones de texte (text boxes) - souvent utilisées dans les schémas
      4. Si le total reste très court : OCR sur les images intégrées
         (cas des documents construits comme des infographies/plans)
    """
    doc = DocxDocument(chemin_fichier)
    morceaux = [p.text for p in doc.paragraphs if p.text.strip()]

    # Texte des tableaux
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    morceaux.append(cell.text)

    # Texte des zones de texte (plans, schémas, infographies)
    zones_texte = extraire_zones_de_texte_docx(chemin_fichier)
    if zones_texte:
        morceaux.append(zones_texte)

    texte = "\n".join(morceaux)

    # Filet de sécurité : si malgré tout on a très peu de texte, le document
    # est probablement construit à partir d'images (plan, pictogrammes) ->
    # on tente l'OCR sur les images intégrées pour ne rien perdre.
    if len(texte.strip()) < SEUIL_TEXTE_MINIMAL_DOCX:
        print(f"   ↳ Peu de texte natif, OCR des images intégrées pour {chemin_fichier.name}...")
        texte_ocr = extraire_images_docx_ocr(chemin_fichier)
        if texte_ocr:
            texte = (texte + "\n" + texte_ocr).strip()

    return texte


def extraire_texte_txt(chemin_fichier: Path) -> str:
    """Lit simplement un fichier texte brut en UTF-8."""
    return chemin_fichier.read_text(encoding="utf-8")


def extraire_texte_xlsx(chemin_fichier: Path) -> str:
    """
    Extrait le contenu textuel d'un fichier Excel (.xlsx), feuille par feuille.
    Chaque ligne devient une ligne de texte, avec les cellules séparées par ' | '.
    Utile pour les fichiers type 'Note de frais' contenant des barèmes ou procédures.
    """
    classeur = openpyxl.load_workbook(chemin_fichier, data_only=True)
    lignes_texte = []

    for nom_feuille in classeur.sheetnames:
        feuille = classeur[nom_feuille]
        lignes_texte.append(f"--- Feuille : {nom_feuille} ---")

        for ligne in feuille.iter_rows(values_only=True):
            cellules = [str(cellule) for cellule in ligne if cellule is not None]
            if cellules:
                lignes_texte.append(" | ".join(cellules))

    return "\n".join(lignes_texte)


def extraire_tous_les_documents(data_dir: Path = DATA_DIR) -> list[dict]:
    """
    Parcourt tous les sous-dossiers de data/ (documents ET formulaires)
    et extrait le texte de chaque fichier trouvé.

    Retourne une liste de dictionnaires {source, text, type, chemin_complet}.
    """
    documents = []

    # Mapping extension -> fonction d'extraction
    extracteurs = {
        ".pdf": extraire_texte_pdf,
        ".docx": extraire_texte_docx,
        ".txt": extraire_texte_txt,
        ".xlsx": extraire_texte_xlsx,
    }

    # On parcourt les deux catégories de dossiers avec leur type associé
    tous_les_dossiers = (
        [(d, "document") for d in DOSSIERS_DOCUMENTS]
        + [(d, "formulaire") for d in DOSSIERS_FORMULAIRES]
    )

    for sous_dossier, type_contenu in tous_les_dossiers:
        dossier = data_dir / sous_dossier
        if not dossier.exists():
            continue

        for fichier in sorted(dossier.iterdir()):
            # Ignorer les fichiers temporaires de verrouillage Word/Office (~$...)
            if fichier.name.startswith("~$"):
                continue

            extension = fichier.suffix.lower()
            if extension not in extracteurs:
                continue

            try:
                texte = extracteurs[extension](fichier)
                if texte.strip():
                    documents.append({
                        "source": fichier.name,
                        "text": texte,
                        "type": type_contenu,
                        "chemin_complet": str(fichier.resolve()),
                    })
                    etiquette = "📎 formulaire" if type_contenu == "formulaire" else "📄 document"
                    print(f" Extrait ({etiquette}) : {fichier.name} ({len(texte)} caractères)")
                else:
                    print(f"  Vide ou illisible : {fichier.name}")
            except Exception as e:
                print(f" Erreur sur {fichier.name} : {e}")

    return documents



if __name__ == "__main__":
    docs = extraire_tous_les_documents()
    print(f"\n Total : {len(docs)} document(s) extrait(s).\n")

    # Aperçu du contenu extrait (les 300 premiers caractères de chaque doc)
    for doc in docs:
        print("=" * 60)
        print(f"Source : {doc['source']}")
        print("-" * 60)
        print(doc["text"][:300] + ("..." if len(doc["text"]) > 300 else ""))
        print()